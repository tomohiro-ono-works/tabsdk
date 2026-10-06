"""定義 DataFrame の新規 Excel ファイル出力。"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .definitions import get_definitions


def write_excel(definitions, output_path: str) -> None:
    target = Path(output_path)
    if target.suffix.lower() != ".xlsx":
        raise ValueError("output_path must end with .xlsx")
    if target.exists():
        raise FileExistsError(f"output already exists: {target}")

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, frame in definitions.items():
        sheet = workbook.create_sheet(name)
        sheet.append(list(frame.columns))
        for values in frame.itertuples(index=False, name=None):
            sheet.append([None if value == "" else value for value in values])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="345D7E")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, column in enumerate(sheet.columns, start=1):
            width = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 12), 60)

    buffer = BytesIO()
    workbook.save(buffer)
    with target.open("xb") as output:
        output.write(buffer.getvalue())


def export_excel(input_path: str, output_path: str) -> None:
    """Workbook を一度だけ解析して新規 `.xlsx` に定義を出力する。"""
    write_excel(get_definitions(input_path), output_path)
